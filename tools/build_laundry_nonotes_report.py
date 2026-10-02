"""Laundry auths that arrived with NO detailed notes (the fixed template note),
split by WellSky status, for one month.
    py tools/build_laundry_nonotes_report.py [--since 2026-09-01] [--until 2026-09-30]
Writes AiHub\\uhc_reviewer_deliverables\\reports\\UHC_Laundry_NoNotes_<since>_<until>.xlsx:
  Summary, Documented (already in WellSky), To Document, one tab per GSSC.
The "Wording" column says whether the note names the GSSC (September wording)
or the Program Manager (team rule 2026-10-02). Prints counts only."""
import argparse, re, sys, requests, importlib.util, pandas as pd
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

R = Path(r"C:\Users\eponce\AiHub\uhc_reviewer"); sys.path.insert(0, str(R))
spec = importlib.util.spec_from_file_location("bridge", R / "uhc_wellsky_journal_bridge.py"); br = importlib.util.module_from_spec(spec); spec.loader.exec_module(br)
import review_notes as rn

ap = argparse.ArgumentParser()
ap.add_argument("--since", default=date.today().strftime("%Y-%m-01"))
ap.add_argument("--until", default=date.today().isoformat())
a = ap.parse_args()
MARK = br.LAUNDRY_NO_DETAILS_MARKER

tok = rn.graph_token(); H = {"Authorization": f"Bearer {tok}"}
site = requests.get(f"{rn.GRAPH}/sites/{rn.SITE_PATH}", headers=H, timeout=60).json()["id"]
web = "https://centralboston.sharepoint.com/sites/DataManagement/Lists/United%20Health%20Care%20Authorizations"
url = (f"{rn.GRAPH}/sites/{site}/lists/{rn.LIST_ID}/items?$expand=fields($select=Title,ClientID,MemberName,PrimaryCareManager,ChangeType,"
       f"JournalNote,CarePlanComments,Services,AuthPeriodStart,AuthPeriodEnd,WellSkyDocumentationStatus)&$top=5000")
best = {}; scanned = 0
while url:
    d = requests.get(url, headers=H, timeout=120).json()
    for it in d.get("value", []):
        f = it.get("fields") or {}; t = (f.get("Title") or "").strip(); created = (it.get("createdDateTime") or "")[:10]
        if not t or created < a.since or created > a.until: continue
        if MARK not in (f.get("JournalNote") or "").lower(): continue
        scanned += 1
        if t not in best or int(it["id"]) < int(best[t][0]["id"]): best[t] = (it, f, created)
    url = d.get("@odata.nextLink")

def wording(note: str) -> str:
    low = (note or "").lower()
    return "Program Manager" if "program manager" in low else ("GSSC" if "gssc" in low else "other")

def row(it, f, created):
    status = f.get("WellSkyDocumentationStatus") or ""
    return {"Received": created, "Auth #": f.get("Title"), "GSSC": f.get("PrimaryCareManager") or "(unassigned)", "Client ID": f.get("ClientID"),
            "Member": f.get("MemberName"), "Change Type": f.get("ChangeType"), "Auth Start": (f.get("AuthPeriodStart") or "")[:10],
            "Auth End": (f.get("AuthPeriodEnd") or "")[:10], "Subject": br.build_subject(f), "Journal Note": f.get("JournalNote"),
            "Service Plan Comments": f.get("CarePlanComments"), "Wording": wording(f.get("JournalNote")),
            "WellSky Status": status or "Not Documented", "SharePoint Item": f"{web}/DispForm.aspx?ID={it['id']}"}

df = pd.DataFrame([row(*v) for v in best.values()]).sort_values(["WellSky Status", "Received", "Auth #"]) if best else pd.DataFrame()
documented = df[df["WellSky Status"] == "Documented"] if len(df) else df
todo = df[df["WellSky Status"] != "Documented"] if len(df) else df
out_dir = Path(r"C:\Users\eponce\AiHub\uhc_reviewer_deliverables\reports"); out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / f"UHC_Laundry_NoNotes_{a.since}_{a.until}.xlsx"

def counts(sub, label):
    return pd.DataFrame({"": [label, "  says GSSC", "  says Program Manager"],
                         "Value": [len(sub), int((sub["Wording"] == "GSSC").sum()) if len(sub) else 0, int((sub["Wording"] == "Program Manager").sum()) if len(sub) else 0]})
summary = pd.concat([pd.DataFrame({"": ["Period", "Laundry auths with no detailed notes (unique auth #)", "List items scanned"], "Value": [f"{a.since} to {a.until}", len(df), scanned]}),
                     counts(documented, "Already documented in WellSky"), counts(todo, "Still to document")], ignore_index=True)
by_gssc = df.groupby(["GSSC", "WellSky Status"]).size().rename("Auths").reset_index() if len(df) else pd.DataFrame()

widths = {"Journal Note": 75, "Service Plan Comments": 45, "Subject": 34, "Member": 22, "GSSC": 22, "Auth #": 13, "Wording": 16, "WellSky Status": 16, "SharePoint Item": 12}
def style(ws, link_col="SharePoint Item"):
    hdr = [c.value for c in ws[1]]
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E79")
    if link_col in hdr:
        li = hdr.index(link_col) + 1
        for r in range(2, ws.max_row + 1):
            c_ = ws.cell(row=r, column=li); c_.hyperlink = c_.value; c_.value = "Open"; c_.font = Font(color="0563C1", underline="single")
    for i, h in enumerate(hdr, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
        for r in range(2, ws.max_row + 1): ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

with pd.ExcelWriter(out, engine="openpyxl") as xw:
    summary.to_excel(xw, sheet_name="Summary", index=False, startrow=0)
    if len(by_gssc): by_gssc.to_excel(xw, sheet_name="Summary", index=False, startrow=len(summary) + 3)
    documented.to_excel(xw, sheet_name="Documented", index=False)
    todo.to_excel(xw, sheet_name="To Document", index=False)
    for g, sub in (df.groupby("GSSC") if len(df) else []):
        sub.to_excel(xw, sheet_name=re.sub(r"[\[\]\*\?/\\:]", " ", g.replace(" (GSSC)", "").replace(" (CM)", "")).strip()[:31] or "unassigned", index=False)
wb = load_workbook(out)
for ws in wb.worksheets:
    if ws.title == "Summary":
        for c in ws[1]: c.font = Font(bold=True)
        ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 24
    else:
        style(ws)
wb.save(out)
print(f"laundry no-notes auths {a.since}..{a.until}: {len(df)} (documented {len(documented)}, to document {len(todo)}); "
      f"wording GSSC={int((df['Wording']=='GSSC').sum()) if len(df) else 0} Program Manager={int((df['Wording']=='Program Manager').sum()) if len(df) else 0}")
print("wrote", out)
