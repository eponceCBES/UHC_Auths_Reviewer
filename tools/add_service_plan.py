r"""Add each consumer's CURRENT WellSky service plan (from HAR) to the laundry
report, as two columns after "Service Plan Comments":

    Current Service Plan (HAR)   the Active care plan: program, dates, care
                                 manager, then one line per active service
                                 allocation (service, units x frequency,
                                 provider, dates)
    Laundry on Plan              Yes / No -- is a laundry service on that plan

Applied to every tab of reports\UHC_Laundry_Auths_by_GSSC.xlsx, matched on
Client ID, saved in place. Rows are PHI: this prints counts only.

"Current" = CARE_PLAN_STATUS = 'Active' (newest start date if several) and
SERVICE_ALLOCATION_STATUS = 'Active' with no end date or an end date in the
future. HAR is a nightly copy of WellSky, so "current" means as of the last
HAR refresh (printed).

    py -3.13 tools/add_service_plan.py [path-to-xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, r"C:\Users\eponce\AiHub\har_powerbi")
import har_pbi  # noqa: E402

REPORT = Path(r"C:\Users\eponce\AiHub\uhc_reviewer_deliverables\reports"
              r"\UHC_Laundry_Auths_by_GSSC.xlsx")
AFTER = "Service Plan Comments"
COL_PLAN = "Current Service Plan (HAR)"
COL_LAUNDRY = "Laundry on Plan"
AGENCY = "Central Boston Elder Services, Inc."

SQL = """
SELECT c.CLIENT_ID,
       p.CARE_PLAN_UUID, p.CARE_PROGRAM_NAME, p.CARE_PLAN_STATUS,
       p.CARE_PLAN_START_DATE, p.CARE_PLAN_END_DATE, p.CARE_PLAN_CARE_MANAGER_NAME,
       p.CARE_PLAN_CARE_MANAGER_IS_PRIMARY, p.CARE_PLAN_LUPDATE_DATETIME,
       p.SERVICE, p.SUBSERVICE, p.SERVICE_CATEGORY, p.PROVIDER,
       p.SERVICE_ALLOCATION_STATUS, p.SERVICE_ALLOCATION_START_DATE,
       p.SERVICE_ALLOCATION_END_DATE, p.UNITS_ALLOCATED, p.FREQUENCY,
       p.ALLOCATION_TYPE, p.SERVICE_ALLOCATION_LUPDATE_DATETIME
FROM HAR_SERVICE_PLANS p
JOIN HAR_CONSUMERS c ON c.CONSUMER_UUID = p.CONSUMER_UUID
WHERE p.AGENCY = '{agency}'
  AND p.CARE_PLAN_STATUS = 'Active'
  AND c.CLIENT_ID IN ({ids})
"""


def _dt(v):
    """HAR dates arrive as text like '9/11/2026 5:07:50 PM' (or blank)."""
    return pd.to_datetime(v, errors="coerce", format="mixed")


def _d(v) -> str:
    t = _dt(v)
    return "" if pd.isna(t) else t.strftime("%m/%d/%Y")


def _fmt_alloc(a: pd.Series) -> str:
    svc = str(a["SERVICE"] or "").strip()
    sub = str(a.get("SUBSERVICE") or "").strip()
    if sub and sub.lower() != svc.lower():
        svc = f"{svc} / {sub}"
    units = a.get("UNITS_ALLOCATED")
    try:
        units = f"{float(units):g}"
    except (TypeError, ValueError):
        units = str(units or "").strip()
    freq = str(a.get("FREQUENCY") or "").strip()
    kind = str(a.get("ALLOCATION_TYPE") or "").strip().lower()
    how = f"{units} x {freq}" if freq else units
    if kind and kind not in ("careplan",):
        how += f" {kind}"
    prov = str(a.get("PROVIDER") or "").strip()
    dates = _d(a.get("SERVICE_ALLOCATION_START_DATE"))
    end = _d(a.get("SERVICE_ALLOCATION_END_DATE"))
    dates = f"{dates} - {end}" if end else (f"from {dates}" if dates else "")
    bits = [svc, how.strip()]
    if prov:
        bits.append(f"via {prov}")
    if dates:
        bits.append(f"({dates})")
    return "  - " + " | ".join(b for b in bits if b)


def summarize(rows: pd.DataFrame) -> tuple[str, str]:
    """(plan text, laundry Yes/No) for one client's rows."""
    if rows.empty:
        return "No ACTIVE care plan in HAR", "No"
    rows = rows.copy()
    rows["_start"] = _dt(rows["CARE_PLAN_START_DATE"])
    newest = rows.sort_values("_start", ascending=False)["CARE_PLAN_UUID"].iloc[0]
    plan = rows[rows["CARE_PLAN_UUID"] == newest]
    h = plan.iloc[0]
    cm = str(h.get("CARE_PLAN_CARE_MANAGER_NAME") or "").strip()
    head = (f"{h['CARE_PROGRAM_NAME']} | Active {_d(h['CARE_PLAN_START_DATE'])}"
            f" - {_d(h['CARE_PLAN_END_DATE']) or 'open'}"
            + (f" | CM: {cm}" if cm else ""))

    today = pd.Timestamp.today().normalize()
    end = _dt(plan["SERVICE_ALLOCATION_END_DATE"])
    live = plan[(plan["SERVICE_ALLOCATION_STATUS"].astype(str).str.strip() == "Active")
                & (end.isna() | (end >= today))]
    live = live.drop_duplicates(subset=["SERVICE", "SUBSERVICE", "PROVIDER",
                                        "UNITS_ALLOCATED", "FREQUENCY",
                                        "SERVICE_ALLOCATION_START_DATE"])
    live = live.sort_values(["SERVICE_CATEGORY", "SERVICE"], na_position="last")
    lines = [_fmt_alloc(a) for _, a in live.iterrows()]
    if not lines:
        lines = ["  - (no active service allocations)"]
    text = "\n".join([head] + lines)
    hay = " ".join(str(v) for v in live[["SERVICE", "SUBSERVICE", "SERVICE_CATEGORY"]]
                   .fillna("").values.ravel()).lower()
    return text, ("Yes" if "laundry" in hay else "No")


def main(path: Path) -> int:
    ids = pd.read_excel(path, sheet_name=0)["Client ID"].dropna().astype(int).unique()
    print(f"report: {path.name}  ({len(ids)} distinct Client IDs)")

    sql = SQL.format(agency=AGENCY, ids=",".join(str(i) for i in ids))
    har = har_pbi.query(sql)
    har["CLIENT_ID"] = pd.to_numeric(har["CLIENT_ID"], errors="coerce").astype("Int64")
    fresh = _dt(pd.concat([har["CARE_PLAN_LUPDATE_DATETIME"],
                           har["SERVICE_ALLOCATION_LUPDATE_DATETIME"]])).max()
    print(f"HAR: {len(har):,} plan/allocation rows for {har['CLIENT_ID'].nunique()} "
          f"clients; newest update in HAR {_d(fresh)}")

    plans: dict[int, tuple[str, str]] = {}
    for cid in ids:
        plans[int(cid)] = summarize(har[har["CLIENT_ID"] == int(cid)])
    n_plan = sum(1 for t, _ in plans.values() if not t.startswith("No ACTIVE"))
    n_laun = sum(1 for _, y in plans.values() if y == "Yes")
    print(f"clients with an Active plan: {n_plan}/{len(ids)}; laundry on plan: {n_laun}")

    wb = load_workbook(path)
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        if COL_PLAN in hdr:                       # re-run: replace, don't duplicate
            for name in (COL_LAUNDRY, COL_PLAN):
                ws.delete_cols(hdr.index(name) + 1)
                hdr = [c.value for c in ws[1]]
        if AFTER not in hdr or "Client ID" not in hdr:
            print(f"  skip tab {ws.title!r}: expected columns missing")
            continue
        at = hdr.index(AFTER) + 2
        cid_col = hdr.index("Client ID") + 1
        # openpyxl's insert_cols moves cells but NOT hyperlinks or column
        # widths -- capture both by header and put them back afterwards.
        links = {}
        for c0 in range(at, len(hdr) + 1):
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, c0)
                if cell.hyperlink is not None:
                    links[(r, c0)] = cell.hyperlink.target
        widths = {h: ws.column_dimensions[get_column_letter(i + 1)].width
                  for i, h in enumerate(hdr)}
        ws.insert_cols(at, 2)
        for (r, c0), target in links.items():
            ws.cell(r, c0 + 2).hyperlink = target
        for i, h in enumerate(hdr):
            if i + 1 >= at and widths.get(h):
                ws.column_dimensions[get_column_letter(i + 3)].width = widths[h]
        for k, name in enumerate((COL_PLAN, COL_LAUNDRY)):
            h = ws.cell(1, at + k, name)
            h.font = Font(bold=True, color="FFFFFF")
            h.fill = PatternFill("solid", fgColor="1F4E79")
        for r in range(2, ws.max_row + 1):
            cid = ws.cell(r, cid_col).value
            try:
                text, yes = plans.get(int(cid), ("", ""))
            except (TypeError, ValueError):
                text, yes = "", ""
            ws.cell(r, at, text).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, at + 1, yes).alignment = Alignment(vertical="top",
                                                          horizontal="center")
        ws.column_dimensions[get_column_letter(at)].width = 70
        ws.column_dimensions[get_column_letter(at + 1)].width = 12
        # widths to the right shifted by two: re-apply from the old letters
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    print(f"saved -> {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1]) if len(sys.argv) > 1 else REPORT))
